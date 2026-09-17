import json
import re
import sys
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import requests
from openpyxl import load_workbook


SOURCE_URL = (
    "https://old.tsetmc.com/tsev2/excel/MarketWatchPlus.aspx?d=0"
)

OUTPUT_PATH = Path("data/assets.json")

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/140.0 Safari/537.36"
    ),
    "Accept": (
        "application/vnd.openxmlformats-officedocument."
        "spreadsheetml.sheet,application/octet-stream,*/*"
    ),
    "Referer": "https://www.tsetmc.com/",
    "Origin": "https://www.tsetmc.com",
}


def normalize_text(value):
    if value is None:
        return ""

    text = str(value).strip()

    # یکسان‌سازی حروف فارسی/عربی
    text = text.replace("ي", "ی")
    text = text.replace("ك", "ک")
    text = text.replace("\u200c", " ")

    # حذف فاصله‌های اضافی
    text = re.sub(r"\s+", " ", text)

    return text


def normalize_symbol(value):
    return normalize_text(value)


def find_column(headers, candidates):
    normalized_headers = [
        normalize_text(header)
        for header in headers
    ]

    for candidate in candidates:
        candidate = normalize_text(candidate)

        for index, header in enumerate(normalized_headers):
            if (
                header == candidate
                or candidate in header
            ):
                return index

    return None


def download_source():
    print("Downloading TSETMC market watch...")

    response = requests.get(
        SOURCE_URL,
        headers=HEADERS,
        timeout=60,
    )

    response.raise_for_status()

    content = response.content

    # XLSX فایل ZIP است و با PK شروع می‌شود.
    if not content.startswith(b"PK"):
        preview = content[:200].decode(
            "utf-8",
            errors="replace",
        )

        raise RuntimeError(
            "TSETMC did not return an Excel file.\n"
            f"Response preview: {preview!r}"
        )

    return content


def parse_market_watch(content):
    print("Parsing market watch...")

    workbook = load_workbook(
        filename=BytesIO(content),
        read_only=True,
        data_only=True,
    )

    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)

        headers = None

        # پیدا کردن ردیف Header
        for row_number, row in enumerate(rows):
            values = list(row)

            joined = " | ".join(
                normalize_text(value)
                for value in values
            )

            if (
                "نماد" in joined
                and "نام" in joined
            ):
                headers = values
                break

            # جلوگیری از جستجوی بی‌نهایت
            if row_number >= 20:
                break

        if headers is None:
            raise RuntimeError(
                "Could not find the Excel header row."
            )

        symbol_column = find_column(
            headers,
            ["نماد"],
        )

        company_column = find_column(
            headers,
            ["نام شرکت", "نام"],
        )

        if symbol_column is None:
            raise RuntimeError(
                "Column 'نماد' was not found."
            )

        if company_column is None:
            raise RuntimeError(
                "Company name column was not found."
            )

        instruments = []
        seen_symbols = set()

        # ادامه خواندن از ردیف بعد از Header
        for row in rows:
            values = list(row)

            if symbol_column >= len(values):
                continue

            if company_column >= len(values):
                continue

            symbol = normalize_symbol(
                values[symbol_column]
            )

            company_name = normalize_text(
                values[company_column]
            )

            # ردیف ناقص
            if not symbol or not company_name:
                continue

            # ردیف Header یا داده غیرنمادی
            if symbol in {"نماد", "سمبل"}:
                continue

            # نمادهای غیرعادی
            if len(symbol) > 40:
                continue

            # حذف تکراری‌ها
            key = symbol.upper()

            if key in seen_symbols:
                continue

            seen_symbols.add(key)

            instruments.append(
                {
                    "symbol": symbol,
                    "company_name": company_name,
                }
            )

        return instruments

    finally:
        workbook.close()


def validate(instruments):
    print(
        f"Validating {len(instruments)} instruments..."
    )

    # اگر تعداد خیلی کم باشد، احتمالاً پاسخ خراب بوده
    if len(instruments) < 1000:
        raise RuntimeError(
            "Safety check failed: parsed fewer than 1000 instruments."
        )

    symbols = [
        item["symbol"]
        for item in instruments
    ]

    # بررسی تکراری نبودن نمادها
    if len(symbols) != len(set(symbols)):
        raise RuntimeError(
            "Safety check failed: duplicate symbols detected."
        )

    # چند نماد پایه فقط برای کنترل سلامت عمومی
    known_symbols = {
        "کچاد",
        "کگل",
        "فملی",
        "فولاد",
    }

    found_known = (
        known_symbols.intersection(set(symbols))
    )

    if len(found_known) == 0:
        raise RuntimeError(
            "Safety check failed: none of the known "
            "reference symbols were found."
        )

    print(
        "Validation passed."
    )


def write_output(instruments):
    OUTPUT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    data = {
        "schema_version": 1,
        "source": SOURCE_URL,
        "generated_at": datetime.now(
            timezone.utc
        ).isoformat(),
        "instrument_count": len(instruments),
        "instruments": sorted(
            instruments,
            key=lambda item: item["symbol"],
        ),
    }

    OUTPUT_PATH.write_text(
        json.dumps(
            data,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    print(
        f"Successfully wrote {OUTPUT_PATH}"
    )


def main():
    content = download_source()

    instruments = parse_market_watch(
        content
    )

    print(
        f"Parsed {len(instruments)} instruments."
    )

    validate(instruments)

    write_output(instruments)

    print(
        "Master instrument list update completed successfully."
    )


if __name__ == "__main__":
    try:
        main()

    except Exception as exc:
        print(
            f"ERROR: {exc}",
            file=sys.stderr,
        )

        sys.exit(1)
